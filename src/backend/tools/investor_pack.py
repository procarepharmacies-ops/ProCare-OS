"""Investor report (Arabic, YoY) + employee motivation deck (Arabic PPTX).

Usage (from src/backend, against the DB ProCare is configured for):
    python tools/investor_pack.py --year 2026 --month 9 [--branch-id N]
        [--target-growth-pct X] [--out DIR] [--no-pdf]

Writes, READ-ONLY on the DB:
    investor_report_YYYY-MM.html   professional Arabic report, month vs same
                                   month last year (YoY) and vs previous month
    investor_report_YYYY-MM.pdf    same, rendered with headless Chromium when found
    team_deck_YYYY-MM.pptx         Arabic RTL motivation deck for the staff,
                                   aimed at next month's results

Every number comes from the DB through monthly_report.collect(); nothing is
estimated. The October target is DERIVED, not invented: last year's same-month
revenue x (1 + growth), where growth = --target-growth-pct if given, else the
month's own YoY growth (floored at 0). The slide states this formula.
Walk-in/credit definition is monthly_report's (see its docstring).
"""
from __future__ import annotations

import argparse
import html
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import monthly_report as mr  # noqa: E402
from app.db.base import SessionLocal  # noqa: E402

TEAL, NAVY, ORANGE = "0D7377", "143B5A", "E67E22"
MONTHS_AR = ["", "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو",
             "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
LOGO = Path(__file__).resolve().parents[2] / "frontend" / "public" / "logo.png"


def pct(cur: float, prev: float) -> float | None:
    return None if not prev else (cur - prev) / prev * 100


def fpct(v: float | None) -> str:
    return "—" if v is None else f"{v:+.1f}%"


def next_month(year: int, month: int) -> tuple[int, int]:
    return (year + (month == 12), 1 if month == 12 else month + 1)


def gather(year: int, month: int, branch_id: int | None, target_growth: float | None) -> dict:
    ny, nm = next_month(year, month)
    with SessionLocal() as s:
        cur = mr.collect(s, year, month, branch_id)
        ly = mr.collect(s, year - 1, month, branch_id)
        nxt_ly = mr.collect(s, ny - 1, nm, branch_id)  # same month as the target month, last year
    k, y = cur["kpis"], ly["kpis"]
    growth = pct(k["net_revenue"], y["net_revenue"])
    if target_growth is None:
        target_growth = max(growth, 0.0) if growth is not None else None
    base = nxt_ly["kpis"]["net_revenue"]
    target = base * (1 + target_growth / 100) if (target_growth is not None and base) else None
    return {"cur": cur, "ly": ly, "k": k, "y": y, "growth": growth, "p": cur["prev"],
            "target": target, "target_growth": target_growth, "target_base": base,
            "next_month": nm, "next_year": ny,
            "has_ly": bool(y["bills"]), "year": year, "month": month}


# ----------------------------- investor report ------------------------------
CSS = f"""
@page{{size:A4;margin:14mm}}
body{{font-family:Cairo,'Segoe UI',Tahoma,'DejaVu Sans',Arial,sans-serif;color:#1d2b36;margin:0;line-height:1.7}}
.cover{{background:#{NAVY};color:#fff;padding:36px 32px;border-bottom:6px solid #{ORANGE}}}
.cover h1{{margin:0;font-size:30px}} .cover p{{margin:6px 0 0;opacity:.85}}
.cover img{{height:56px;float:left;background:#fff;border-radius:10px;padding:4px}}
main{{padding:18px 28px}} h2{{color:#{TEAL};border-bottom:2px solid #{TEAL};padding-bottom:4px;margin-top:26px}}
.kpis{{display:flex;flex-wrap:wrap;gap:10px}}
.kpi{{flex:1 1 30%;background:#eef6f6;border-right:5px solid #{TEAL};border-radius:8px;padding:10px 14px}}
.kpi small{{display:block;color:#51636f}} .kpi b{{font-size:22px;color:#{NAVY}}}
.up{{color:#1e8449;font-weight:700}} .down{{color:#c0392b;font-weight:700}}
table{{border-collapse:collapse;width:100%;margin:8px 0 14px;font-size:14px}}
th{{background:#{NAVY};color:#fff}} th,td{{border:1px solid #cfd8dc;padding:6px 9px;text-align:right}}
td.n{{direction:ltr;text-align:left}} .note{{background:#fff7ec;border-right:5px solid #{ORANGE};padding:8px 12px;font-size:13px}}
.foot{{color:#6b7c88;font-size:12px;margin-top:22px}}
"""


def _delta_cell(cur: float, prev: float, invert: bool = False) -> str:
    d = pct(cur, prev)
    if d is None:
        return "<td class=n>—</td>"
    good = d < 0 if invert else d > 0
    return f'<td class="n {"up" if good else "down"}">{d:+.1f}%</td>'


def _chart(cur_daily: dict, ly_daily: dict, month_days: int) -> str:
    w, h, pad = 760, 220, 28
    top = max(list(cur_daily.values()) + list(ly_daily.values()) + [1.0])
    bw = (w - 2 * pad) / month_days
    bars = "".join(
        f'<rect x="{pad + (d - 1) * bw + 1:.1f}" y="{h - pad - v / top * (h - 2 * pad):.1f}" '
        f'width="{bw - 2:.1f}" height="{v / top * (h - 2 * pad):.1f}" fill="#{TEAL}"/>'
        for d, v in sorted(cur_daily.items()))
    pts = " ".join(f"{pad + (d - 0.5) * bw:.1f},{h - pad - v / top * (h - 2 * pad):.1f}"
                   for d, v in sorted(ly_daily.items()))
    line = f'<polyline points="{pts}" fill="none" stroke="#{ORANGE}" stroke-width="2"/>' if pts else ""
    ticks = "".join(f'<text x="{pad + (d - 0.5) * bw:.1f}" y="{h - 10}" font-size="9" text-anchor="middle">{d}</text>'
                    for d in range(1, month_days + 1, 3))
    return (f'<svg viewBox="0 0 {w} {h}" width="100%" role="img" aria-label="المبيعات اليومية">{bars}{line}{ticks}</svg>'
            f'<p class=foot>الأعمدة: السنة الحالية — الخط البرتقالي: نفس الشهر العام الماضي</p>')


def render_report(g: dict) -> str:
    k, y, p, mo, yr = g["k"], g["y"], g["p"], g["month"], g["year"]
    name = MONTHS_AR[mo]
    rows = [("إجمالي المبيعات", "revenue", False), ("صافي المبيعات (بعد المرتجعات)", "net_revenue", False),
            ("إجمالي الربح", "gross_profit", False), ("عدد الفواتير", "bills", False),
            ("متوسط الفاتورة", "avg_bill", False), ("المرتجعات", "returns", True),
            ("الخصومات", "discount", True), ("المبيعات الآجلة", "credit_value", False)]
    ky = dict(y, gross_profit=y["revenue"] - y.get("cogs", 0)) if "cogs" in y else y
    body_rows = ""
    for label, key, inv in rows:
        cv = k.get(key, 0.0)
        pv_prev = p.get(key, 0.0) if key in p else 0.0
        lv = ky.get(key, 0.0) if g["has_ly"] and key in ky else 0.0
        body_rows += (f"<tr><td>{label}</td><td class=n>{mr._fmt(float(cv))}</td>"
                      f"<td class=n>{mr._fmt(float(lv)) if g['has_ly'] and key in ky else '—'}</td>"
                      f"{_delta_cell(cv, lv, inv) if g['has_ly'] and key in ky else '<td class=n>—</td>'}"
                      f"<td class=n>{mr._fmt(float(pv_prev)) if key in p else '—'}</td>"
                      f"{_delta_cell(cv, pv_prev, inv) if key in p else '<td class=n>—</td>'}</tr>")
    credit_share = k["credit_value"] / k["revenue"] * 100 if k["revenue"] else 0.0
    ly_note = "" if g["has_ly"] else (
        '<p class=note>لا توجد بيانات لنفس الشهر من العام الماضي في قاعدة البيانات؛ '
        "تم إخفاء مقارنة العام على العام.</p>")
    branches = "".join(f"<tr><td>{html.escape(n)}</td><td class=n>{mr._fmt(v['revenue'])}</td>"
                       f"<td class=n>{v['bills']}</td></tr>" for n, v in g["cur"]["branches"].items())
    prods = "".join(f"<tr><td>{html.escape(r['name'])}</td><td class=n>{mr._fmt(r['units'])}</td>"
                    f"<td class=n>{mr._fmt(r['revenue'])}</td></tr>" for r in g["cur"]["top_products"][:10])
    logo = '<img src="logo.png" alt="">' if LOGO.exists() else ""
    growth = g["growth"]
    headline = (f"صافي مبيعات {name} {yr}: {mr._fmt(k['net_revenue'])} ج.م"
                + (f" — {fpct(growth)} مقارنةً بـ {name} {yr - 1}" if growth is not None else ""))
    return f"""<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">
<title>تقرير المستثمرين — {name} {yr}</title><style>{CSS}</style></head><body>
<div class="cover">{logo}<h1>صيدليات بروكير — تقرير الأداء الشهري للمستثمرين</h1>
<p>{name} {yr} · مقارنة سنوية (YoY) ومقارنة بالشهر السابق</p></div><main>
<h2>الملخص التنفيذي</h2><p><b>{html.escape(headline)}</b></p>
<div class="kpis">
<div class="kpi"><small>صافي المبيعات</small><b>{mr._fmt(k['net_revenue'])}</b><small>YoY {fpct(growth)}</small></div>
<div class="kpi"><small>عدد الفواتير</small><b>{k['bills']:,}</b><small>YoY {fpct(pct(k['bills'], y['bills'])) if g['has_ly'] else '—'}</small></div>
<div class="kpi"><small>هامش الربح الإجمالي</small><b>{k['margin_pct']:.1f}%</b><small>إجمالي الربح {mr._fmt(k['gross_profit'])}</small></div>
</div>{ly_note}
<h2>مؤشرات الأداء — {name} {yr} مقابل {name} {yr - 1}</h2>
<table><tr><th>المؤشر</th><th>{name} {yr}</th><th>{name} {yr - 1}</th><th>التغير السنوي</th><th>الشهر السابق</th><th>التغير الشهري</th></tr>{body_rows}</table>
<h2>المبيعات اليومية</h2>{_chart(g['cur']['daily'], g['ly']['daily'], 31 if mo in (1,3,5,7,8,10,12) else 30 if mo != 2 else 29)}
<h2>الأداء حسب الفرع</h2><table><tr><th>الفرع</th><th>المبيعات</th><th>الفواتير</th></tr>{branches}</table>
<h2>أعلى ١٠ أصناف مبيعًا</h2><table><tr><th>الصنف</th><th>الكمية</th><th>القيمة</th></tr>{prods}</table>
<h2>التعرض الائتماني (البيع الآجل)</h2>
<p>بلغت المبيعات الآجلة {mr._fmt(k['credit_value'])} ج.م عبر {k['credit_bills']} فاتورة، أي {credit_share:.1f}% من إجمالي المبيعات. التفاصيل في تقرير الفواتير الآجلة المنفصل.</p>
<h2>منهجية وملاحظات</h2><p class=note>المصدر: قاعدة بيانات ProCare المنقولة من eStock (قراءة فقط). إجمالي الربح = المبيعات − (الكمية × سعر الشراء).
الآجل = فاتورة غير مرتجعة لعميل مسجّل والمدفوع (نقدي + فيزا) أقل من الصافي؛ يُراجَع مع كشوف حسابات العملاء. لا توجد تقديرات في هذا التقرير.</p>
<p class=foot>أُعدّ آليًا بواسطة ProCare OS.</p></main></body></html>"""


# --------------------------------- deck -------------------------------------
def build_deck(g: dict, path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    def rgb(h: str) -> RGBColor:
        return RGBColor.from_string(h)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    k, mo, yr = g["k"], g["month"], g["year"]
    name, nname = MONTHS_AR[mo], MONTHS_AR[g["next_month"]]

    def text(slide, x, y, w, h, s, size=24, bold=False, color="143B5A", align=PP_ALIGN.RIGHT):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(s.split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para._p.get_or_add_pPr().set("rtl", "1")
            r = para.add_run()
            r.text = line
            r.font.size, r.font.bold = Pt(size), bold
            r.font.color.rgb = rgb(color)
            r.font.name = "Tahoma"
            rpr = r._r.get_or_add_rPr()
            cs = rpr.makeelement(qn("a:cs"), {"typeface": "Tahoma"})
            rpr.append(cs)
        return tb

    def box(slide, x, y, w, h, fill):
        shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        shp.line.fill.background()
        return shp

    def base(title: str, dark: bool = False):
        s = prs.slides.add_slide(blank)
        box(s, 0, 0, 13.333, 7.5, NAVY if dark else "FFFFFF")
        if not dark:
            box(s, 0, 0, 13.333, 1.2, NAVY)
            box(s, 0, 1.2, 13.333, 0.08, ORANGE)
            text(s, 0.5, 0.15, 12.3, 0.9, title, 34, True, "FFFFFF")
        if LOGO.exists():
            s.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.2) if not dark else Inches(0.4), height=Inches(0.8))
        return s

    def stat(slide, x, y, label, value, sub=""):
        box(slide, x, y, 3.9, 2.0, "EEF6F6")
        box(slide, x + 3.82, y, 0.08, 2.0, TEAL)
        text(slide, x + 0.1, y + 0.05, 3.6, 0.5, label, 18, False, "51636F")
        text(slide, x + 0.1, y + 0.55, 3.6, 0.9, value, 28, True, NAVY)
        if sub:
            text(slide, x + 0.1, y + 1.4, 3.6, 0.5, sub, 16, True, TEAL)

    # 1 — title
    s = base("", dark=True)
    text(s, 0.8, 2.0, 11.7, 1.4, f"شكرًا لكل فريق بروكير", 54, True, "FFFFFF")
    text(s, 0.8, 3.5, 11.7, 1.0, f"نتائج {name} {yr} .. وهدفنا في {nname}", 30, False, "F5B041")
    text(s, 0.8, 5.6, 11.7, 0.6, "صيدليات بروكير — معًا نصنع الفرق", 20, False, "BFD8D8")

    # 2 — what we achieved
    s = base(f"إنجازنا في {name} {yr}")
    growth = g["growth"]
    stat(s, 8.9, 1.8, "صافي المبيعات", f"{mr._fmt(k['net_revenue'])}", f"{fpct(growth)} عن {name} {yr - 1}" if g["has_ly"] else "")
    stat(s, 4.7, 1.8, "عدد الفواتير", f"{k['bills']:,}", f"{fpct(pct(k['bills'], g['y']['bills']))} عن العام الماضي" if g["has_ly"] else "")
    stat(s, 0.5, 1.8, "متوسط الفاتورة", f"{mr._fmt(k['avg_bill'])}", "")
    text(s, 0.6, 4.3, 12.2, 2.4,
         "كل فاتورة هي مريض اطمأن على دوائه، وكل ابتسامة خدمة هي سبب ثقة العملاء فينا.\n"
         "هذه الأرقام صنعتموها أنتم — الصيادلة والكاشير وفريق المخزن والدعم.", 24, False, NAVY)

    # 3 — YoY
    s = base(f"{name} {yr} مقابل {name} {yr - 1}")
    if g["has_ly"]:
        y = g["y"]
        items = [("صافي المبيعات", k["net_revenue"], y["net_revenue"]), ("عدد الفواتير", k["bills"], y["bills"]),
                 ("متوسط الفاتورة", k["avg_bill"], y["avg_bill"])]
        for i, (lab, c, p_) in enumerate(items):
            d = pct(c, p_)
            stat(s, 8.9 - i * 4.2, 1.8, lab, fpct(d), f"{mr._fmt(float(c))} مقابل {mr._fmt(float(p_))}")
        msg = ("نمو حقيقي مقارنةً بالعام الماضي — نفخر بكم!" if (growth or 0) > 0
               else "العام الماضي كان قويًا — وفرصتنا الآن لنتفوق عليه معًا.")
        text(s, 0.6, 4.4, 12.2, 1.6, msg, 28, True, TEAL)
    else:
        text(s, 0.6, 2.5, 12.2, 2, "لا تتوافر بيانات نفس الشهر من العام الماضي للمقارنة.", 26, False, NAVY)

    # 4 — October target
    s = base(f"هدفنا في {nname}")
    if g["target"]:
        text(s, 0.6, 1.7, 12.2, 1.0, f"{mr._fmt(g['target'])} ج.م صافي مبيعات", 44, True, ORANGE)
        text(s, 0.6, 2.9, 12.2, 1.2,
             f"المعادلة: صافي مبيعات {nname} {g['next_year'] - 1} ({mr._fmt(g['target_base'])}) × (1 + {g['target_growth']:.1f}%)", 20, False, "51636F")
    else:
        text(s, 0.6, 2.2, 12.2, 1.6, "سنحدد هدف الشهر بعد اعتماد الإدارة.", 32, True, ORANGE)
    text(s, 0.6, 4.3, 12.2, 2.4,
         "كيف نصل إليه؟\n• ابتسامة وترحيب مع كل عميل\n• التأكد من توافر الصنف أو اقتراح البديل المناسب\n• متابعة العملاء الدائمين والآجل بانتظام", 24, False, NAVY)

    # 5 — close
    s = base("", dark=True)
    text(s, 0.8, 2.3, 11.7, 1.6, "نجاحنا يبدأ من كل واحد فينا", 48, True, "FFFFFF")
    text(s, 0.8, 4.2, 11.7, 1.0, f"هيا نجعل {nname} أفضل شهر!", 32, True, "F5B041")
    prs.save(str(path))


# ---------------------------------- main ------------------------------------
def render_pdf(html_path: Path, pdf_path: Path) -> bool:
    chrome = (shutil.which("chromium") or shutil.which("google-chrome") or shutil.which("chrome")
              or next(iter(sorted(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux/chrome"))), None)
              or next((p for p in (Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
                                   Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")) if p.exists()), None))
    if not chrome:
        return False
    r = subprocess.run([str(chrome), "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()], capture_output=True, timeout=120)
    return r.returncode == 0 and pdf_path.exists()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--month", type=int, required=True)
    ap.add_argument("--branch-id", type=int, default=None)
    ap.add_argument("--target-growth-pct", type=float, default=None)
    ap.add_argument("--out", default=".tmp/reports")
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    if not 1 <= a.month <= 12:
        print("ERROR: --month must be 1..12", file=sys.stderr)
        return 2
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    g = gather(a.year, a.month, a.branch_id, a.target_growth_pct)
    tag = f"{a.year}-{a.month:02d}"
    hp = out / f"investor_report_{tag}.html"
    hp.write_text(render_report(g), encoding="utf-8")
    if LOGO.exists():
        shutil.copyfile(LOGO, out / "logo.png")
    build_deck(g, out / f"team_deck_{tag}.pptx")
    pdf = "skipped" if a.no_pdf else ("ok" if render_pdf(hp, out / f"investor_report_{tag}.pdf") else "chromium not found")
    print(f"OK {tag}: bills={g['k']['bills']} YoY={fpct(g['growth'])} pdf={pdf} -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
